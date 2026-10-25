"""Build catalog.json products from unitprice.xlsx.

Smart Store columns: 품명, 옵션, 세부옵션, 설정가격, 옵션가.
옵션가 is a delta added to 설정가격. 소스 and 아이스팩 become add-ons.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / 'unitprice.xlsx'
CATALOG = ROOT / 'catalog.json'
R2 = 'https://pub-236307febd5647328aae37b69df5373d.r2.dev'

SHEET_CATEGORY = {
    '소포장 대용량': 1,
    '닭안심가슴살 단품': 1,
    '특수부위': 2,
    '윙,봉,닭다리살': 3,
    '닭한마리&오리고기': 4,
    '소고기': 5,
    '돼지고기': 6,
}

# Match longer / more specific titles first. id None => new product.
RULES = [
    ('1kgx10p 소포장', '닭안심·가슴살 소포장', f'{R2}/products/chicken-breast-premium.webp', None),
    ('닭안심살 1kg단품', '신선 닭안심살', f'{R2}/products/chicken-breast-premium.webp', None),
    ('닭가슴살 1kg단품', '신선 닭가슴살', f'{R2}/products/chicken-breast-premium.webp', 14),
    ('무뼈닭발 500g', '무뼈닭발', f'{R2}/products/boneless-feet-premium.webp', 1),
    ('육수용 닭뼈', '육수용 닭뼈', f'{R2}/products/bones-premium.webp', 2),
    ('닭어깨살', '무염지 닭어깨살', f'{R2}/products/chicken-shoulder-premium.webp', 4),
    ('닭가슴연골', '닭가슴연골·오돌뼈', f'{R2}/products/cartilage-premium.webp', 11),
    ('닭염통', '닭염통·간·알집', f'{R2}/products/organs-premium.webp', 5),
    ('닭근위', '닭근위', f'{R2}/products/gizzard-premium.webp', 3),
    ('닭목살', '닭목살', f'{R2}/products/neck-premium.webp', 6),
    ('안창살', '닭 안창살', f'{R2}/products/inner-premium.webp', 12),
    ('손질뼈닭발', '손질 뼈닭발', f'{R2}/products/feet-premium.webp', 13),
    ('닭껍질 10kg', '닭껍질 대용량', f'{R2}/products/skin-premium.webp', 8),
    ('무뼈닭발 냉동 2kg', '무뼈닭발 대용량', f'{R2}/products/boneless-feet-premium.webp', 7),
    ('생닭껍질', '생 닭껍질', f'{R2}/products/skin-premium.webp', 9),
    ('뼈없는 닭다리살 통정육', '국내산 닭다리살', f'{R2}/products/domestic-thigh-skin-on-premium.webp', None),
    ('순살통닭다리살', '브라질 순살 닭다리살', f'{R2}/products/boneless-thigh-skin-on-premium.webp', None),
    ('냉동 닭봉', '닭봉·윙·통날개', f'{R2}/products/wings-mix-premium.webp', None),
    ('통정육 2kg 5팩', '브라질 통정육 대용량', f'{R2}/products/bulk-thigh-skin-on-premium.webp', None),
    ('장각5조각', '냉장 장각 5조각', f'{R2}/products/whole-leg-five-premium.webp', None),
    ('생닭다리(북채)', '닭다리(북채)', f'{R2}/products/chicken-drumsticks-premium.webp', None),
    ('장각4조각', '냉장 장각 4조각', f'{R2}/products/whole-leg-four-premium.webp', None),
    ('오리안심 1kg', '생 오리안심', f'{R2}/categories/whole.png', None),
    ('오리훈제', '오리훈제', f'{R2}/categories/whole.png', None),
    ('통마리 생닭', '통마리 생닭', f'{R2}/categories/whole.png', None),
    ('닭도리육', '닭도리육', f'{R2}/categories/whole.png', None),
    ('삼계탕용', '삼계탕용 생닭', f'{R2}/categories/whole.png', None),
    ('오리안심 대용량', '오리안심 대용량', f'{R2}/categories/whole.png', None),
    ('소대패목살', '소대패목살', f'{R2}/categories/beef.png', None),
    ('우삼겹', '우삼겹 돌돌이', f'{R2}/categories/beef.png', None),
    ('목전지', '목전지', f'{R2}/categories/pork.png', None),
    ('대패삼겹살', '대패삼겹살', f'{R2}/categories/pork.png', None),
    ('대패목살말이', '대패목살말이', f'{R2}/categories/pork.png', None),
]

FALLBACK_BASE = (
    ('무뼈닭발', '2kg', 22000),
)


def clean(value) -> str:
    if value is None:
        return ''
    return re.sub(r'\s+', ' ', str(value).replace('\\', '')).strip()


def parse_money(value):
    if value is None or value == '':
        return None
    if isinstance(value, (int, float)):
        return int(value)
    text = clean(value).replace(',', '').replace('원', '')
    if not text:
        return None
    return int(text)


def is_marker(name: str) -> bool:
    if not name or name == '품명':
        return True
    if name.startswith('*'):
        return True
    if '아이스팩' in name:
        return True
    return False


def is_sauce_row(opt1: str, opt2: str, group: str) -> bool:
    if opt1 == '소스' or group == '소스' and not opt1:
        return bool(opt2 or opt1)
    return False


def match_rule(title: str):
    for needle, name, image, product_id in RULES:
        if needle in title.replace(' ', ''):
            return name, image, product_id
        if needle in title:
            return name, image, product_id
    return None


def load_rows():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    parsed = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header = [clean(c) for c in next(ws.iter_rows(values_only=True))]
        simple = '옵션' not in ''.join(header)
        category_id = SHEET_CATEGORY[sheet_name]
        current = None
        group = ''
        for raw in ws.iter_rows(min_row=2, values_only=True):
            cells = list(raw) + [None] * (5 - len(raw))
            name = clean(cells[0])
            if simple:
                if not name or is_marker(name):
                    if current and '아이스팩' in name:
                        current['addons'].setdefault('아이스팩', 500)
                    continue
                rule = match_rule(name)
                if rule is None:
                    raise SystemExit(f'이름 규칙 없음: {name}')
                short, image, product_id = rule
                price = parse_money(cells[1])
                if not price or price <= 0:
                    raise SystemExit(f'단가 없음: {name}')
                current = {
                    'title': name,
                    'name': short,
                    'image': image,
                    'match_id': product_id,
                    'category_id': category_id,
                    'base': price,
                    'options': [{'label': '기본', 'price': price, 'delta': 0}],
                    'addons': {},
                }
                parsed.append(current)
                continue

            if name and not is_marker(name):
                rule = match_rule(name)
                if rule is None:
                    raise SystemExit(f'이름 규칙 없음: {name}')
                short, image, product_id = rule
                base = parse_money(cells[3])
                if not base or base <= 0:
                    for a, b, fallback in FALLBACK_BASE:
                        if a in name and b in name:
                            base = fallback
                            break
                if not base or base <= 0:
                    raise SystemExit(f'설정가격 없음: {name} row={cells}')
                current = {
                    'title': name,
                    'name': short,
                    'image': image,
                    'match_id': product_id,
                    'category_id': category_id,
                    'base': base,
                    'options': [],
                    'addons': {},
                }
                parsed.append(current)
                group = ''

            if current is None:
                continue
            if '아이스팩' in name:
                current['addons'].setdefault('아이스팩', 500)

            opt1 = clean(cells[1])
            opt2 = clean(cells[2])
            delta = parse_money(cells[4])
            listed = parse_money(cells[3])
            if opt1:
                group = opt1
            if is_sauce_row(opt1, opt2, group):
                label = opt2 or opt1
                label = label.replace('소스', '소스') if False else label
                if label and label != '소스':
                    price = delta if delta is not None else (listed or 0)
                    if price < 0:
                        raise SystemExit(f'소스 가격 오류: {current["name"]} {label}')
                    current['addons'].setdefault(label, price)
                continue
            if opt1 and '약재' in opt1 and not opt2:
                price = listed if listed is not None else (delta or 0)
                current['addons'].setdefault(opt1, price)
                continue
            if not opt1 and not opt2:
                continue
            if opt1 and opt2:
                label = f'{opt1} / {opt2}'
            elif opt2:
                label = f'{group} / {opt2}' if group and group not in opt2 else opt2
            else:
                label = opt1
            label = clean(label)
            if not label or label == '소스':
                continue
            price = current['base'] + (delta or 0)
            if price < 0:
                raise SystemExit(f'옵션가 음수: {current["name"]} {label} {price}')
            if any(opt['label'] == label for opt in current['options']):
                continue
            current['options'].append({'label': label, 'price': price, 'delta': delta or 0})
        for product in parsed:
            if product['category_id'] == category_id and not product['options']:
                product['options'].append({'label': '기본', 'price': product['base'], 'delta': 0})
    return parsed


def order_options(product):
    options = product['options']
    base_hits = [opt for opt in options if opt['price'] == product['base']]
    if not base_hits:
        return options
    first = base_hits[0]
    rest = [opt for opt in options if opt is not first]
    return [first, *rest]


def build_products(parsed):
    used_ids = set()
    products = []
    seen_titles = set()
    next_id = 15
    for item in parsed:
        if item['title'] in seen_titles:
            continue
        seen_titles.add(item['title'])
        product_id = item['match_id']
        if product_id in used_ids:
            continue
        if product_id is None:
            while next_id in used_ids or next_id == 10:
                next_id += 1
            product_id = next_id
            next_id += 1
        used_ids.add(product_id)
        options = [{'label': opt['label'], 'price': opt['price']} for opt in order_options(item)]
        addons = [{'label': label, 'price': price} for label, price in item['addons'].items()]
        products.append(
            {
                'id': product_id,
                'category_id': item['category_id'],
                'name': item['name'],
                'option': options[0]['label'],
                'price': options[0]['price'],
                'image': item['image'],
                'options': options,
                'addons': addons,
                '_title': item['title'],
            }
        )
    products.sort(key=lambda row: row['id'])
    return products


def main():
    parsed = load_rows()
    products = build_products(parsed)
    catalog = json.loads(CATALOG.read_text(encoding='utf-8'))
    kept = [row for row in catalog['products'] if row['id'] == 10]
    if not kept:
        raise SystemExit('기존 업소용 무뼈닭발(id 10)을 찾지 못했습니다.')
    public = []
    for row in products:
        row = dict(row)
        row.pop('_title', None)
        public.append(row)
    public.extend(kept)
    public.sort(key=lambda row: row['id'])
    catalog['products'] = public
    CATALOG.write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'products {len(public)}')
    for row in public:
        opts = row.get('options') or []
        adds = row.get('addons') or []
        print(f"{row['id']:3} cat={row['category_id']} {row['price']:>7} opts={len(opts):2} addons={len(adds):2} {row['name']}")


if __name__ == '__main__':
    main()
